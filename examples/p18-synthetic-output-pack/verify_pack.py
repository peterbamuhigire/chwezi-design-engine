"""Check the P18 synthetic pack's editable files, cached arithmetic, and exports."""

from __future__ import annotations

import json
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parent
RENDERS = ROOT / "renders" / "final"
EXPECTED = {"synthetic-internal-proposal": 1, "synthetic-cash-timing-report": 2,
            "synthetic-cash-timing-model": 3, "synthetic-executive-briefing": 5}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def main() -> None:
    data = json.loads((ROOT / "source-data.json").read_text(encoding="utf-8"))
    for name in ("synthetic-internal-proposal.docx", "synthetic-cash-timing-report.docx"):
        doc = Document(ROOT / name)
        require(bool(doc.paragraphs), f"{name}: document has no paragraphs")
        require(bool(doc.tables), f"{name}: expected semantic data table")
    report = Document(ROOT / "synthetic-cash-timing-report.docx")
    report_text = "\n".join(p.text for p in report.paragraphs)
    require("Illustrative" in report_text and "translation not reviewed" in report_text.lower(),
            "report must retain synthetic and translation boundaries")

    book = openpyxl.load_workbook(ROOT / "synthetic-cash-timing-model.xlsx", data_only=False)
    values = openpyxl.load_workbook(ROOT / "synthetic-cash-timing-model.xlsx", data_only=True)
    require({"Summary", "Inputs", "Cash bridges"}.issubset(book.sheetnames), "workbook sheets missing")
    expected = [(0, 0, 0, 0), (-60, 60, 0, 100), (-120, 120, 45, 200)]
    for row, (trough, need, gap, receivables) in enumerate(expected, start=4):
        actual = tuple(values["Summary"].cell(row, col).value for col in (4, 5, 7, 8))
        require(actual == (trough, need, gap, receivables), f"workbook row {row}: {actual}")
        require(book["Summary"].cell(row, 4).data_type == "f", f"summary row {row} lost formulas")
    require(data["monthly_billings"] * data["periods"] == 400, "fixture source changed unexpectedly")
    summary_chart = book["Summary"]._charts[0]
    require(len(summary_chart.ser) == 1 and summary_chart.type == "col",
            "workbook visual must chart one gap measure across the three cases")

    deck = Presentation(ROOT / "synthetic-executive-briefing.pptx")
    require(len(deck.slides) == 5, "deck slide count changed")
    slide_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    require("45 SCU" in slide_text and "NOT ASSESSED" in slide_text, "deck conclusion/boundaries missing")

    reopened = RENDERS / "reopened"
    require(len(Document(reopened / "synthetic-internal-proposal.docx").tables) > 0, "DOCX reopen failed")
    require(len(Document(reopened / "synthetic-cash-timing-report.docx").tables) >= 3, "report reopen failed")
    reopened_book = openpyxl.load_workbook(reopened / "synthetic-cash-timing-model.xlsx", data_only=False)
    require(reopened_book["Summary"]["D4"].data_type == "f", "XLSX roundtrip removed formulas")
    require(len(Presentation(reopened / "synthetic-executive-briefing.pptx").slides) == 5,
            "PPTX reopen failed")

    for stem, page_count in EXPECTED.items():
        pdf = fitz.open(RENDERS / f"{stem}.pdf")
        require(len(pdf) == page_count, f"{stem}: expected {page_count} PDF pages, got {len(pdf)}")
        require(all(page.get_text().strip() for page in pdf), f"{stem}: blank rendered page")
        for index in range(page_count):
            png = RENDERS / "pages" / f"{stem}-{index + 1:02}.png"
            require(png.is_file() and png.stat().st_size > 5000, f"missing/empty page render: {png.name}")
    expected_pages = {f"{stem}-{index:02}.png" for stem, count in EXPECTED.items()
                      for index in range(1, count + 1)}
    expected_pages.add("cash-timing-chart-mobile-preview.png")
    actual_pages = {path.name for path in (RENDERS / "pages").glob("*.png")}
    require(actual_pages == expected_pages, f"unexpected or missing final previews: {sorted(actual_pages ^ expected_pages)}")
    phone_preview = RENDERS / "pages" / "cash-timing-chart-mobile-preview.png"
    require(phone_preview.is_file(), "390-pixel infographic preview missing")
    print("P18 synthetic pack verification: PASS")
    print("editable formats=DOCX(2), XLSX(1), PPTX(1); PDFs=1/2/3/5 pages; formula cache and roundtrip checks passed")


if __name__ == "__main__":
    main()
