"""Check propagation of the P18 one-input cash-buffer change."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import fitz
import openpyxl
from pptx import Presentation


ROOT = Path(__file__).resolve().parent / "renders" / "final" / "mutation-proof"


def require(ok: bool, message: str) -> None:
    if not ok:
        raise AssertionError(message)


def main() -> None:
    data = json.loads((ROOT / "source-data.json").read_text(encoding="utf-8"))
    require(data["available_liquidity"] == 100, "mutated buffer input missing")
    book = openpyxl.load_workbook(ROOT / "recalculated" / "synthetic-cash-timing-model.xlsx", data_only=True)
    require([book["Summary"].cell(row, 7).value for row in (4, 5, 6)] == [0, 0, 20],
            "recalculated XLSX gaps should be 0/0/20")
    require(book["Inputs"]["B7"].value == 100, "recalculated XLSX buffer should be 100")

    report_pdf = fitz.open(ROOT / "synthetic-cash-timing-report.pdf")
    report_text = "\n".join(page.get_text() for page in report_pdf)
    require("20 SCU" in report_text and "100 SCU" in report_text, "report did not carry changed gap/buffer")
    require("45 SCU" not in report_text and "75 SCU" not in report_text, "report retained stale baseline figures")

    deck_pdf = fitz.open(ROOT / "synthetic-executive-briefing.pdf")
    deck_text = "\n".join(page.get_text() for page in deck_pdf)
    require("20 SCU" in deck_text and "100 SCU" in deck_text, "deck did not carry changed gap/buffer")
    require("45 SCU" not in deck_text and "75 SCU" not in deck_text, "deck retained stale baseline figures")
    deck = Presentation(ROOT / "synthetic-executive-briefing.pptx")
    require(len(deck.slides) == 5, "mutated deck did not preserve slide structure")
    with zipfile.ZipFile(ROOT / "synthetic-executive-briefing.pptx") as package:
        descriptions = " ".join(package.read(name).decode("utf-8") for name in package.namelist()
                                 if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
    require("20 SCU" in descriptions and "100 SCU" in descriptions,
            "PPTX image descriptions did not carry changed values")
    with zipfile.ZipFile(ROOT / "synthetic-cash-timing-report.docx") as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    require("20 SCU" in document_xml and "100 SCU" in document_xml,
            "DOCX text/figure description did not carry changed values")
    print("P18 one-input mutation propagation: PASS (buffer 75 -> 100; gap 45 -> 20 across XLSX, DOCX and PPTX)")


if __name__ == "__main__":
    main()
