"""Rebuild the P18 synthetic DOCX/XLSX/PPTX format demonstration."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.dimensions import ColumnDimension
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor as SlideRGB
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as SInches, Pt as SPt


ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "source-data.json").read_text(encoding="utf-8"))
INK = "18313A"
TEAL = "176B68"
PAPER = "F7F5F0"
MUTED = "58676D"
RULE = "D9E1DF"
BLUE_INPUT = "1466A5"
RED = "9C3F35"
DISPLAY = "Spectral"
BODY = "Atkinson Hyperlegible"
OUT = {
    "chart": ROOT / "cash-timing-chart.png",
    "proposal": ROOT / "synthetic-internal-proposal.docx",
    "report": ROOT / "synthetic-cash-timing-report.docx",
    "workbook": ROOT / "synthetic-cash-timing-model.xlsx",
    "deck": ROOT / "synthetic-executive-briefing.pptx",
}


def simulate(lag: int) -> dict:
    billings = [DATA["monthly_billings"]] * DATA["periods"]
    receipts = [0] * DATA["periods"]
    for month, amount in enumerate(billings):
        if month + lag < DATA["periods"]:
            receipts[month + lag] += amount
    cash = []
    balance = DATA["opening_cash"]
    for receipt in receipts:
        balance += receipt - DATA["monthly_cash_costs"]
        cash.append(balance)
    trough = min([DATA["opening_cash"], *cash])
    need = max(0, -trough)
    return {
        "billings": billings,
        "receipts": receipts,
        "costs": [DATA["monthly_cash_costs"]] * DATA["periods"],
        "cash": cash,
        "trough": trough,
        "need": need,
        "liquidity": DATA["available_liquidity"],
        "gap": max(0, need - DATA["available_liquidity"]),
        "receivables": sum(billings) - sum(receipts),
    }


RESULTS = [(scenario, simulate(scenario["lag"])) for scenario in DATA["scenarios"]]
assert [x[1]["need"] for x in RESULTS] == [0, 60, 120]
assert [x[1]["gap"] for x in RESULTS] == [0, 0, 45]
assert all(sum(x[1]["billings"]) == 400 for x in RESULTS)


def set_cell_fill(cell, colour: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), colour)
    tc_pr.append(shd)


def mark_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def setup_doc(path: Path, title: str, subtitle: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.72)
    sec.left_margin, sec.right_margin = Inches(.86), Inches(.86)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = BODY, Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name in ("Caption", "Footer", "List Paragraph", "List Bullet", "List Number"):
        if name in doc.styles:
            doc.styles[name].font.name = BODY
            doc.styles[name].font.color.rgb = RGBColor.from_string(MUTED if name in ("Caption", "Footer") else INK)
    for name, size, colour in (("Title", 27, INK), ("Heading 1", 17, TEAL), ("Heading 2", 12, INK)):
        style = doc.styles[name]
        style.font.name, style.font.size = DISPLAY if name != "Heading 2" else BODY, Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = "CHWEZI  /  SYNTHETIC FORMAT REVIEW"
    header.style = doc.styles["Caption"]
    header.runs[0].font.name = BODY
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("INTERNAL DEMONSTRATION  •  NOT CLIENT MATERIAL     |     ")
    run.font.name, run.font.size = BODY, Pt(7)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title)
    r.font.name, r.font.size, r.font.bold = DISPLAY, Pt(26), True
    r.font.color.rgb = RGBColor.from_string(INK)
    sub = doc.add_paragraph(subtitle)
    sub.paragraph_format.space_after = Pt(14)
    sub.runs[0].font.name, sub.runs[0].font.size = BODY, Pt(12)
    sub.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    badge = doc.add_paragraph(DATA["label"])
    badge.paragraph_format.space_after = Pt(20)
    badge.runs[0].font.bold = True
    badge.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    return doc


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    mark_repeat_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_fill(cell, INK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name, r.font.size, r.font.bold = BODY, Pt(8), True
                r.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = str(text)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_fill(cells[idx], PAPER)
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name, r.font.size = BODY, Pt(7.7)
                    r.font.color.rgb = RGBColor.from_string(INK)
        if widths:
            for idx, width in enumerate(widths):
                cells[idx].width = Inches(width)
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    return table


def make_chart() -> None:
    from PIL import Image, ImageDraw, ImageFont

    design = Path(__file__).resolve().parents[2]
    display_path = design / "fonts/01-formal-institutional/spectral/Spectral-Regular.ttf"
    body_path = design / "fonts/08-body-ui-workhorses/public-sans/PublicSans-VF.ttf"
    img = Image.new("RGB", (1600, 900), "#F7F5F0")
    draw = ImageDraw.Draw(img)
    title = ImageFont.truetype(str(display_path), 62) if display_path.exists() else ImageFont.load_default()
    body = ImageFont.truetype(str(body_path), 48) if body_path.exists() else ImageFont.load_default()
    if hasattr(body, "set_variation_by_axes"):
        body.set_variation_by_axes([500])
    label = ImageFont.truetype(str(body_path), 44) if body_path.exists() else ImageFont.load_default()
    if hasattr(label, "set_variation_by_axes"):
        label.set_variation_by_axes([650])
    kicker = ImageFont.truetype(str(body_path), 30) if body_path.exists() else ImageFont.load_default()
    if hasattr(kicker, "set_variation_by_axes"):
        kicker.set_variation_by_axes([600])
    metric = ImageFont.truetype(str(display_path), 132) if display_path.exists() else ImageFont.load_default()
    draw.text((82, 46), "SYNTHETIC CASH UNITS", font=kicker, fill=f"#{MUTED}")
    draw.text((80, 104), "Only the two-period delay leaves a gap", font=title, fill=f"#{INK}")
    draw.text((84, 194), "Unfunded amount beyond the assumed 75 SCU buffer", font=body, fill=f"#{MUTED}")
    colours = ["#176B68", "#1466A5", "#9C3F35"]
    names = ["No lag", "One period", "Two periods"]
    gaps = [result["gap"] for _, result in RESULTS]
    for x, name, gap, colour in zip((78, 570, 1062), names, gaps, colours):
        draw.rounded_rectangle((x, 306, x + 460, 785), radius=28, fill="#FFFFFF", outline=f"#{RULE}", width=3)
        draw.rounded_rectangle((x, 306, x + 460, 324), radius=8, fill=colour)
        draw.text((x + 36, 363), name.upper(), font=label, fill=f"#{INK}")
        value = str(gap)
        bbox = draw.textbbox((0, 0), value, font=metric)
        value_width = bbox[2] - bbox[0]
        draw.text((x + (460 - value_width) // 2, 418), value, font=metric, fill=colour)
        draw.text((x + 150, 622), "SCU gap", font=body, fill=f"#{INK}")
    img.save(OUT["chart"], dpi=(300, 300), optimize=True)


def make_report() -> None:
    doc = setup_doc(OUT["report"], "Collection delay and cash timing", "Synthetic research note  /  Four-period operating cash bridge")
    doc.add_heading("Finding", 1)
    doc.add_paragraph("With billings and cash costs held constant, longer collection delays deepen the temporary cash trough. In this artificial four-period example, the two-period delay creates a 45 SCU gap beyond the assumed 75 SCU buffer. The numbers demonstrate timing arithmetic only.", style="Normal")
    doc.add_paragraph("The result is conditional on the input fixture. It does not establish a customer, market, revenue-recognition rule, financing source, accounting basis, tax treatment or rollout decision.")
    doc.add_heading("Scenario comparison", 1)
    summary = [[s["id"], str(s["lag"]), str(sum(r["billings"])), str(r["trough"]), str(r["need"]), str(r["gap"]), str(r["receivables"])] for s, r in RESULTS]
    add_table(doc, ["Scenario", "Lag", "Billings", "Cash trough", "Need", "Gap", "Ending AR"], summary, [1.45, .45, .7, .85, .6, .6, .75])
    note = doc.add_paragraph("1  SCU is a synthetic cash unit. Ending receivables equal total fixture billings less receipts within the four periods; this is not an accounting receivables balance.")
    note.paragraph_format.space_before = Pt(3)
    note.runs[0].font.size = Pt(8)
    note.runs[0].font.italic = True
    doc.add_heading("Buffer gap by collection delay", 1)
    pic = doc.add_picture(str(OUT["chart"]), width=Inches(5.2))
    doc_pr = pic._inline.docPr
    doc_pr.set("descr", "Three synthetic collection-delay cases show the gap beyond an assumed 75 SCU buffer: zero for no lag, zero for one-period lag, and 45 SCU for two-period lag.")
    cap = doc.add_paragraph("Figure 1. Only the two-period delay leaves a gap beyond the assumed buffer.")
    cap.style = doc.styles["Caption"]
    source = doc.add_paragraph(f"Source: {DATA['source_id']} (synthetic fixture; {DATA['assumption_date']}).")
    source.style = doc.styles["Caption"]
    doc.add_page_break()
    doc.add_heading("Reperformance detail", 1)
    doc.add_paragraph("The next table is intentionally long enough to exercise row continuity and repeated headers across a page boundary. Every value traces to the one synthetic fixture; calculations are reproduced in the editable workbook.")
    rows = []
    for scenario, result in RESULTS:
        for month in range(DATA["periods"]):
            rows.append([scenario["id"], str(month + 1), str(result["billings"][month]), str(result["receipts"][month]), str(result["costs"][month]), str(result["cash"][month])])
    add_table(doc, ["Scenario", "Period", "Billings", "Receipts", "Cash costs", "Closing cash"], rows, [1.45, .55, .82, .82, .82, 1.0])
    note = doc.add_paragraph("Source: synthetic fixture only. No source transaction, accounting period or report balance is represented.")
    note.runs[0].font.size = Pt(8)
    doc.add_heading("Inputs and interpretation", 1)
    inputs = [
        ["Billings / cash costs", "100 / 60 SCU per period", "Artificial inputs; held constant across all scenarios."],
        ["Opening cash / buffer", "0 / 75 SCU", "The buffer is assumed; it is not committed funding."],
        ["Horizon / collection", "4 periods / 100%", "Simplifying test assumptions; credit risk is not modelled."],
        ["Reporting framework", "Not selected", "No entity-specific financial statements are produced."],
        ["Revenue / tax", "NOT ASSESSED", "No recognition, tax rate or statutory treatment is included."],
        ["Source / date", f"{DATA['source_id']} / {DATA['assumption_date']}", "Synthetic fixture; date is not an effective date for law or policy."],
        ["Reconciliation / approval", "Arithmetic only / none", "No subledger, control account, report, Controller or tax sign-off was supplied."],
        ["Translation / portability", "French sample / local PDF", "Translation review, font embedding and recipient-machine proof remain NOT ASSESSED."],
    ]
    add_table(doc, ["Item", "Value", "Interpretation / limit"], inputs, [1.35, 1.4, 3.6])
    doc.add_heading("Bilingual disclosure sample", 1)
    doc.add_paragraph("Illustrative scenario only; these figures do not describe a customer or forecast.")
    doc.add_paragraph("Scénario illustratif uniquement ; ces chiffres ne décrivent pas un client ni une prévision. (French layout sample; translation not reviewed.)")
    doc.add_heading("Countercase and next evidence", 1)
    doc.add_paragraph("A longer delay matters only if the inputs and decision horizon are representative. The same gap can disappear with opening liquidity, lower costs, partial collections or a different horizon. None of those alternatives was tested here.")
    doc.add_paragraph("Before a real decision, obtain an approved transaction and reporting scope, trace source-event timing, confirm available funding, and have the accountable finance reviewer reperform the calculation. Do not convert this example into a forecast or funding recommendation.")
    doc.save(OUT["report"])


def make_proposal() -> None:
    doc = setup_doc(OUT["proposal"], "Proposal: keep the output test", "Internal decision brief  /  No fee, client or delivery commitment")
    doc.add_heading("Decision requested", 1)
    doc.add_paragraph("Keep this pack as a synthetic, local production test while domain and recipient review gates remain open. It is not a client deliverable or a request to publish or spend.")
    doc.add_heading("Why this test", 1)
    doc.add_paragraph("The same limited collection-timing example is carried into four editable formats. That makes it possible to check whether figures, caveats and table structure survive document, workbook, slide and PDF production.")
    doc.add_heading("What the pack demonstrates", 1)
    add_table(doc, ["Output", "Test performed", "Limit"], [
        ["Research report DOCX/PDF", "Source note, conclusion limits, chart, long repeated-header table and bilingual sample.", "No external research or human translation review."],
        ["Cash timing XLSX/PDF", "Editable assumptions, linked formulas, summary, chart and print view.", "Not a business plan, forecast or financial statement."],
        ["Executive PPTX/PDF", "Conclusion titles, comparison chart, action boundaries and slide export.", "No real sponsor brief or live-room test."],
    ], [1.5, 2.65, 2.2])
    doc.add_heading("Evidence and objection", 1)
    doc.add_paragraph("The synthetic calculation is independently reproducible from the four-period input table; billings and costs remain fixed across scenarios. A reasonable objection is that good local rendering does not prove compatibility with every customer's Office version, printer, font set or accessibility workflow. That objection is valid, so this output is evidence of one local toolchain only.")
    doc.add_heading("Boundaries", 1)
    for item in DATA["limits"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Recommended next step", 1)
    doc.add_paragraph("Retain the generated source files and render manifest for internal format testing. Obtain a named language reviewer, domain owner and recipient-application check before any client-facing use. No budget or timeline is proposed because no sponsor scope was supplied.")
    doc.save(OUT["proposal"])


def make_workbook() -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    inputs = wb.create_sheet("Inputs")
    bridge = wb.create_sheet("Cash bridges")
    chartdata = wb.create_sheet("Chart data")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        ws.page_margins.left = ws.page_margins.right = .35
        ws.page_margins.top = ws.page_margins.bottom = .45
    inputs.append(["Synthetic collection-timing inputs"])
    inputs.append(["Field", "Value", "Unit / boundary"])
    for row in [
        ["Periods", DATA["periods"], "periods"],
        ["Opening cash", DATA["opening_cash"], "SCU"],
        ["Monthly billings", DATA["monthly_billings"], "SCU/period"],
        ["Monthly cash costs", DATA["monthly_cash_costs"], "SCU/period"],
        ["Available liquidity", DATA["available_liquidity"], "SCU; assumed buffer, not committed funding"],
        ["Collection rate", DATA["collection_rate_percent"], "%; simplified test assumption"],
        ["Scenario: no lag", 0, "periods"],
        ["Scenario: one-period lag", 1, "periods"],
        ["Scenario: two-period lag", 2, "periods"],
    ]:
        inputs.append(row)
    inputs.append(["Scope", "SYNTHETIC ONLY", "No client, market, tax or accounting conclusion."])
    inputs.append(["Source", DATA["source_id"], DATA["source_note"]])
    for row in range(3, 13):
        inputs.cell(row, 2).font = Font(name=BODY, color=BLUE_INPUT)
    summary.append(["Collection delay changes temporary cash needs"])
    summary.append(["Synthetic cash timing only  /  SCU is not a real currency"])
    summary.append(["Scenario", "Lag (periods)", "Total billings", "Cash trough", "Funding need", "Assumed buffer", "Unfunded gap", "Ending receivables"])
    for idx, source_row in enumerate((9, 10, 11), start=4):
        summary.cell(idx, 1, f"=Inputs!A{source_row}")
        summary.cell(idx, 2, f"=Inputs!B{source_row}")
        block_start = 4 + (idx - 4) * 5
        summary.cell(idx, 3, f"=SUM('Cash bridges'!C{block_start}:C{block_start+3})")
        summary.cell(idx, 4, f"=MIN(Inputs!B4,'Cash bridges'!G{block_start}:G{block_start+3})")
        summary.cell(idx, 5, f"=MAX(0,-D{idx})")
        summary.cell(idx, 6, "=Inputs!B7")
        summary.cell(idx, 7, f"=MAX(0,E{idx}-F{idx})")
        summary.cell(idx, 8, f"=C{idx}-SUM('Cash bridges'!D{block_start}:D{block_start+3})")
    summary.append(["SYNTHETIC ONLY  /  Arithmetic illustration; see Inputs for scope limits."])
    bridge.append(["Monthly cash bridge by synthetic collection lag"])
    bridge.append(["Scenario", "Period", "Billings", "Cash receipts", "Cash costs", "Opening cash", "Closing cash", "Collection lag"])
    for idx, source_row in enumerate((9, 10, 11)):
        start = 4 + idx * 5
        lag_cell = f"Inputs!B{source_row}"
        for month in range(1, DATA["periods"] + 1):
            r = start + month - 1
            bridge.cell(r, 1, f"=Inputs!A{source_row}")
            bridge.cell(r, 2, month)
            bridge.cell(r, 3, "=Inputs!B5")
            bridge.cell(r, 4, f"=IF(B{r}>{lag_cell},C{r},0)")
            bridge.cell(r, 5, "=Inputs!B6")
            bridge.cell(r, 6, "=Inputs!B4" if month == 1 else f"=G{r-1}")
            bridge.cell(r, 7, f"=F{r}+D{r}-E{r}")
            bridge.cell(r, 8, f"={lag_cell}")
    chartdata.append(["Gap comparison for the summary chart"])
    chartdata.append([])
    chartdata.append(["Collection delay", "Unfunded gap (SCU)"])
    for row_num, (summary_row, scenario) in enumerate(zip(range(4, 7), DATA["scenarios"]), start=4):
        chartdata.cell(row_num, 1, scenario["id"])
        chartdata.cell(row_num, 2, f"=Summary!G{summary_row}")
    for ws, widths in [(summary, [26, 16, 18, 16, 16, 17, 16, 20]), (inputs, [32, 19, 74]), (bridge, [25, 12, 16, 16, 16, 16, 16, 16]), (chartdata, [18, 20, 24, 23])]:
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width
        ws.freeze_panes = "A4" if ws is summary else "A3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = "Synthetic internal example — not client or accounting output"
    header_fill = PatternFill("solid", fgColor=INK)
    accent_fill = PatternFill("solid", fgColor=TEAL)
    for ws, header_row in [(summary, 3), (inputs, 2), (bridge, 2), (chartdata, 3)]:
        for cell in ws[header_row]:
            cell.fill = header_fill
            cell.font = Font(name=BODY, color="FFFFFF", bold=True, size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for ws in wb.worksheets:
        ws.row_dimensions[1].height = 26
        ws["A1"].font = Font(name=DISPLAY, size=16, bold=True, color=INK)
        ws["A1"].alignment = Alignment(vertical="center")
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if c.data_type == "f":
                    c.font = Font(name=BODY, color=INK, size=9)
                if c.column > 1 and c.row > 3 and isinstance(c.value, (int, float)):
                    c.number_format = '#,##0;[Red](#,##0);0'
                    c.alignment = Alignment(horizontal="right", vertical="top")
        ws.sheet_view.zoomScale = 90
    summary["A7"].font = Font(name=BODY, color=RED, bold=True, size=9)
    summary["A7"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    summary.row_dimensions[7].height = 22
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Only the two-period delay leaves a gap"
    chart.x_axis.title, chart.y_axis.title = "Collection delay", "Unfunded gap (SCU)"
    chart.height, chart.width = 8, 22
    chart.add_data(Reference(chartdata, min_col=2, min_row=3, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(chartdata, min_col=1, min_row=4, max_row=6))
    chart.legend = None
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 50
    chart.y_axis.majorUnit = 10
    chart.series[0].graphicalProperties.solidFill = TEAL
    chart.series[0].graphicalProperties.line.solidFill = TEAL
    summary.add_chart(chart, "A8")
    summary.print_area = "A1:H26"
    bridge.print_area = "A1:H18"
    inputs.print_area = "A1:C13"
    chartdata.sheet_state = "hidden"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUT["workbook"])


def add_text(slide, x, y, w, h, text, size=18, colour=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(SInches(x), SInches(y), SInches(w), SInches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = SInches(.06)
    tf.margin_top = tf.margin_bottom = SInches(.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = font, SPt(size), bold
    r.font.color.rgb = SlideRGB.from_string(colour)
    return box


def make_deck() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SInches(13.333), SInches(7.5)
    blank = prs.slide_layouts[6]
    slides = [prs.slides.add_slide(blank) for _ in range(5)]
    for slide in slides:
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = SlideRGB.from_string(PAPER)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, SInches(.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = SlideRGB.from_string(TEAL); bar.line.fill.background()
    add_text(slides[0], .8, 1.05, 11.6, .7, "Collection delay shifts cash needs", 30, INK, True, DISPLAY)
    add_text(slides[0], .85, 2.0, 10.8, .7, "A four-period synthetic fixture keeps billings and cash costs fixed", 18, TEAL, False)
    add_text(slides[0], .85, 3.65, 11.5, .35, "UNFUNDED GAP AFTER THE ASSUMED BUFFER", 11, MUTED, True)
    for i, (name, value, colour) in enumerate((("No lag", "0 SCU", TEAL), ("One-period lag", "0 SCU", BLUE_INPUT), ("Two-period lag", "45 SCU", RED))):
        x = .85 + i * 4.1
        add_text(slides[0], x, 4.12, 3.55, .55, value, 25, colour, True, DISPLAY)
        add_text(slides[0], x, 4.78, 3.55, .38, name, 13, INK, True)
    add_text(slides[0], .85, 5.65, 11.4, .45, DATA["label"], 11, MUTED, True)
    add_text(slides[0], .85, 6.25, 11.4, .5, "SCU is not a real currency. Internal format test; no client or accounting conclusion.", 10, MUTED)
    add_text(slides[1], .75, .65, 11.8, .6, "The scenario holds billings and costs constant", 25, INK, True, DISPLAY)
    for i, (label, value, detail, color) in enumerate([
        ("Billings", "400 SCU", "same in all cases", TEAL),
        ("Cash costs", "240 SCU", "same in all cases", BLUE_INPUT),
        ("Available buffer", "75 SCU", "assumption only", RED),
    ]):
        x = .85 + i * 4.1
        shape = slides[1].shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, SInches(x), SInches(2.05), SInches(3.55), SInches(2.0))
        shape.fill.solid(); shape.fill.fore_color.rgb = SlideRGB.from_string("FFFFFF"); shape.line.color.rgb = SlideRGB.from_string(RULE)
        add_text(slides[1], x+.2, 2.35, 3.2, .6, value, 25, color, True, DISPLAY, PP_ALIGN.CENTER)
        add_text(slides[1], x+.2, 3.0, 3.2, .4, label, 14, INK, True, BODY, PP_ALIGN.CENTER)
        add_text(slides[1], x+.2, 3.45, 3.2, .35, detail, 10, MUTED, False, BODY, PP_ALIGN.CENTER)
    add_text(slides[1], .85, 5.2, 11.6, .8, "Assumptions are artificial test inputs; the fixture does not model revenue recognition, taxes, customer credit risk or actual financing.", 13, MUTED)
    add_text(slides[2], .75, .65, 11.8, .6, "Buffer gap by collection delay", 25, INK, True, DISPLAY)
    pic = slides[2].shapes.add_picture(str(OUT["chart"]), SInches(1.25), SInches(1.48), width=SInches(9.65))
    pic._element._nvXxPr.cNvPr.set("descr", "Three-case synthetic comparison of gap beyond an assumed 75 SCU buffer: 0 SCU with no lag, 0 SCU with a one-period lag, 45 SCU with a two-period lag.")
    add_text(slides[2], .95, 7.15, 11.3, .2, "Synthetic calculation; not a forecast or committed funding gap.", 8, MUTED)
    add_text(slides[3], .75, .65, 11.8, .6, "Only receipt timing changes across the three cases", 25, INK, True, DISPLAY)
    headers = ["Lag", "Cash trough", "Funding need", "Buffer", "Gap", "Ending AR"]
    xs = [.9, 3.0, 5.0, 7.0, 9.0, 10.9]
    for x, h in zip(xs, headers): add_text(slides[3], x, 1.8, 1.7, .55, h, 12, TEAL, True)
    for row, (scenario, result) in enumerate(RESULTS):
        y = 2.55 + row * 1.05
        if row % 2 == 0:
            sh = slides[3].shapes.add_shape(MSO_SHAPE.RECTANGLE, SInches(.8), SInches(y-.12), SInches(11.75), SInches(.8))
            sh.fill.solid(); sh.fill.fore_color.rgb = SlideRGB.from_string("EAF0EE"); sh.line.fill.background()
        values = [f"{scenario['lag']} period(s)", str(result["trough"]), str(result["need"]), str(result["liquidity"]), str(result["gap"]), str(result["receivables"])]
        for x, v in zip(xs, values): add_text(slides[3], x, y, 1.7, .46, v, 13, INK, row == 2)
    add_text(slides[3], .9, 6.15, 11.6, .6, "Receivables are fixture billings not received by period 4; this is not an accounting subledger balance.", 10, MUTED)
    add_text(slides[4], .75, .65, 11.8, .6, "Use this pack to test file production, not to approve a transaction", 24, INK, True, DISPLAY)
    for i, text in enumerate([
        "Reporting entity, framework and transaction source: NOT ASSESSED.",
        "Tax, revenue recognition and statutory presentation: NOT ASSESSED.",
        "Independent Controller, language-owner and client review: NOT ASSESSED.",
        "Before real use, reperform against an approved source contract and reconcile outputs.",
    ]):
        add_text(slides[4], 1.0, 1.75 + i * .9, 11.3, .58, f"{i+1:02d}   {text}", 15, TEAL if i == 3 else INK, i == 3)
    add_text(slides[4], .95, 6.15, 11.3, .48, "No client, market, tax or financing claim is made. Do not publish or use as a client pack.", 10, MUTED)
    prs.save(OUT["deck"])


def main() -> None:
    make_chart()
    make_report()
    make_proposal()
    make_workbook()
    make_deck()
    print(json.dumps({key: str(value) for key, value in OUT.items()}, indent=2))


if __name__ == "__main__":
    main()
